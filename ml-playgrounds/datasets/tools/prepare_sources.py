"""원본을 그대로 쓸 수 없는 자료를 build_dataset.py에 넣을 영어 CSV로 다듬는다.

    python prepare_sources.py <raw 폴더> <clean 폴더> recipes [id ...]

id를 주면 그 자료만 만든다 (예: power-demand). 주지 않으면 전부.

단위 변환(인치 -> cm), 정답을 새는 열 빼기, 목표 열을 맨 끝으로 옮기기,
헤더 없는 UCI 원본에 열 이름 붙이기처럼 레시피만으로는 할 수 없는 일을 한다.
레시피 JSON도 여기서 함께 쓴다 - 버섯처럼 코드표가 큰 것은 손으로 옮기다 틀린다.

raw 폴더에 받아 둘 것 (파일 이름 그대로):
  galton.tsv         https://dataverse.harvard.edu/api/access/datafile/3037856
  diabetes.csv       https://hbiostat.org/data/repo/diabetes.csv
  titanic3.csv       https://hbiostat.org/data/repo/titanic3.csv
  penguins.csv       https://raw.githubusercontent.com/allisonhorst/palmerpenguins/main/inst/extdata/penguins.csv
  zoo/zoo.data       https://archive.ics.uci.edu/static/public/111/zoo.zip 를 zoo/에 푼다
  mushroom/agaricus-lepiota.data
                     https://archive.ics.uci.edu/static/public/73/mushroom.zip 를 mushroom/에 푼다
  ccpp/CCPP/Folds5x2_pp.xlsx
                     https://archive.ics.uci.edu/static/public/294/combined+cycle+power+plant.zip 를 ccpp/에 푼다
  realestate/Real estate valuation data set.xlsx
                     https://archive.ics.uci.edu/static/public/477/real+estate+valuation+data+set.zip 를 realestate/에 푼다
  pm25/PRSA_data_2010.1.1-2014.12.31.csv
                     https://archive.ics.uci.edu/static/public/381/beijing+pm2+5+data.zip 를 pm25/에 푼다
  kpx.csv            https://www.data.go.kr/data/15065266/fileData.do 에서 받은 CSV (CP949)
                     기온은 스크립트가 Open-Meteo에서 직접 받는다 (인터넷 필요)
  fraud/fraudTrain.csv, fraud/fraudTest.csv
                     https://www.kaggle.com/datasets/kartik2112/fraud-detection (로그인 필요, 약 500MB)
  rps.zip            https://storage.googleapis.com/learning-datasets/rps.zip
  rps-test.zip       https://storage.googleapis.com/learning-datasets/rps-test-set.zip

엑셀 원본은 openpyxl, 가위바위보 이미지는 Pillow가 있어야 한다.

그다음 자료마다:
  python build_dataset.py recipes/<id>.json --csv <clean 폴더>/<id>.csv --out <자산 폴더>
레시피의 csv_name이 목록인 자료(신용카드)는 --csv를 그 순서대로 여러 번 준다.
가위바위보는 이미지라 레시피 없이 <clean 폴더>/rock-paper-scissors.zip이 바로 나온다.
"""
import csv, json, os, sys

RAW, CLEAN, RECIPES = sys.argv[1:4]
os.makedirs(CLEAN, exist_ok=True)


def read(path, delim=",", header=True):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f, delimiter=delim) if any(c.strip() for c in r)]
    return (rows[0], rows[1:]) if header else (None, rows)


def write(name, header, rows):
    with open(os.path.join(CLEAN, name), "w", encoding="utf-8", newline="") as f:
        csv.writer(f, lineterminator="\n").writerows([header] + rows)


def recipe(rec):
    with open(os.path.join(RECIPES, rec["id"] + ".json"), "w", encoding="utf-8") as f:
        json.dump(rec, f, ensure_ascii=False, indent=2)
        f.write("\n")


def cm(v):
    return "" if v in ("", "NA") else f"{float(v) * 2.54:.1f}"


def kg(v):
    return "" if v in ("", "NA") else f"{float(v) * 0.45359237:.1f}"


def H(en, ko):
    return {"en": en, "ko": ko}



# ── Galton ────────────────────────────────────────────────────────────────
def galton_heights():
    h, rows = read(os.path.join(RAW, "galton.tsv"), "\t")
    i = {c: h.index(c) for c in h}
    out = [[cm(r[i["father"]]), cm(r[i["mother"]]), r[i["gender"]], r[i["kids"]], cm(r[i["height"]])]
           for r in rows]
    write("galton-heights.csv", ["father (cm)", "mother (cm)", "gender", "kids", "height (cm)"], out)
    recipe({
        "id": "galton-heights",
        "csv_name": "galton-heights.csv",
        "source_url": "https://doi.org/10.7910/DVN/T0HSJ1",
        "license": "CC0 1.0",
        "title": {"ko": "부모와 자녀의 키 (Galton)", "en": "Galton Family Heights"},
        "target": {"ko": "자녀 키 (height, cm)", "en": "height (cm)"},
        "headers": [
            H("father (cm)", "아버지 키 (father, cm)"),
            H("mother (cm)", "어머니 키 (mother, cm)"),
            H("gender", "자녀 성별 (gender)"),
            H("kids", "가족의 자녀 수 (kids)"),
            H("height (cm)", "자녀 키 (height, cm)"),
        ],
        "values": {"gender": {"ko": {"M": "아들 (M)", "F": "딸 (F)"},
                              "en": {"M": "Son (M)", "F": "Daughter (F)"}}},
        "body": {
            "ko": "1880년대 프랜시스 골턴이 모은 가족 키 기록입니다. 아버지와 어머니의 키로 자녀의 키를 예측합니다.\n"
                  "성별 열로 아들과 딸을 나누어 각각 회귀선을 그려 보면 '평균으로의 회귀'가 어디서 나온 말인지 보입니다.",
            "en": "Francis Galton's family height records from the 1880s. Predict a child's height from the father's and mother's.\n"
                  "Split by gender and fit sons and daughters separately to see where 'regression to the mean' got its name.",
        },
        "notes": {
            "ko": ["키를 인치에서 센티미터로 바꾸었습니다 (소수 첫째 자리).",
                   "가족 번호(family)와, 성별을 0/1로 다시 적은 male·female 열을 뺐습니다."],
            "en": ["Heights were converted from inches to centimetres (one decimal place).",
                   "Dropped the family ID and the male/female columns, which repeat gender as 0/1."],
        },
    })


# ── Diabetes (Willems et al., Vanderbilt) ────────────────────────────────
def diabetes():
    h, rows = read(os.path.join(RAW, "diabetes.csv"))
    i = {c: h.index(c) for c in h}
    out = []
    for r in rows:
        g = r[i["glyhb"]]
        if g == "":
            continue
        out.append([
            r[i["gender"]], r[i["age"]], r[i["location"]],
            cm(r[i["height"]]), kg(r[i["weight"]]), r[i["frame"]],
            cm(r[i["waist"]]), cm(r[i["hip"]]),
            r[i["bp.1s"]], r[i["bp.1d"]],
            r[i["chol"]], r[i["hdl"]], r[i["stab.glu"]],
            "yes" if float(g) >= 7.0 else "no",
        ])
    write("diabetes.csv",
          ["gender", "age", "location", "height (cm)", "weight (kg)", "frame",
           "waist (cm)", "hip (cm)", "systolic bp", "diastolic bp",
           "cholesterol", "hdl", "glucose", "diabetes"], out)
    recipe({
        "id": "diabetes",
        "csv_name": "diabetes.csv",
        "source_url": "https://hbiostat.org/data/",
        "license": "Free to use with attribution - \"Data obtained from http://hbiostat.org/data courtesy of the Vanderbilt University Department of Biostatistics.\"",
        "title": {"ko": "당뇨 검사", "en": "Diabetes Screening"},
        "target": {"ko": "당뇨 여부 (diabetes)", "en": "diabetes"},
        "headers": [
            H("gender", "성별 (gender)"),
            H("age", "나이 (age)"),
            H("location", "거주 지역 (location)"),
            H("height (cm)", "키 (height, cm)"),
            H("weight (kg)", "몸무게 (weight, kg)"),
            H("frame", "체격 (frame)"),
            H("waist (cm)", "허리둘레 (waist, cm)"),
            H("hip (cm)", "엉덩이둘레 (hip, cm)"),
            H("systolic bp", "수축기 혈압 (systolic bp)"),
            H("diastolic bp", "이완기 혈압 (diastolic bp)"),
            H("cholesterol", "총콜레스테롤 (cholesterol)"),
            H("hdl", "HDL 콜레스테롤 (hdl)"),
            H("glucose", "혈당 (glucose)"),
            H("diabetes", "당뇨 여부 (diabetes)"),
        ],
        "values": {
            "gender": {"ko": {"male": "남성 (male)", "female": "여성 (female)"}},
            "location": {"ko": {"Buckingham": "버킹엄 (Buckingham)", "Louisa": "루이자 (Louisa)"}},
            "frame": {"ko": {"small": "작음 (small)", "medium": "보통 (medium)", "large": "큼 (large)", "": ""}},
            "diabetes": {"ko": {"yes": "당뇨 (yes)", "no": "정상 (no)"}},
        },
        "body": {
            "ko": "미국 버지니아주 두 지역의 성인 약 400명을 검사한 기록입니다. 키·몸무게·허리둘레·혈압·콜레스테롤·혈당으로 당뇨 여부를 맞힙니다.\n"
                  "당뇨 여부는 당화혈색소(glyhb) 7.0 이상을 기준으로 정했습니다. 그 수치는 정답을 그대로 알려 주므로 파일에서 뺐습니다.",
            "en": "Screening records of about 400 adults from two counties in Virginia, USA. Predict diabetes from height, weight, waist, blood pressure, cholesterol and glucose.\n"
                  "Diabetes is defined as glycosylated haemoglobin (glyhb) of 7.0 or more. That value gives the answer away, so it is not in the file.",
        },
        "notes": {
            "ko": ["당화혈색소(glyhb)로 당뇨 여부 열을 만들고 glyhb 열은 뺐습니다. glyhb가 없는 13행도 뺐습니다.",
                   "키·허리둘레·엉덩이둘레는 인치에서 센티미터로, 몸무게는 파운드에서 킬로그램으로 바꾸었습니다.",
                   "대부분 비어 있는 두 번째 혈압 측정값, 피험자 번호, 콜레스테롤/HDL 비율(두 열로 계산되는 값), 채혈 시각 열을 뺐습니다.",
                   "원 연구: Willems JP et al. (1997) Prevalence of coronary heart disease risk factors among rural blacks. Journal of Community Health 22:365-375."],
            "en": ["Built the diabetes column from glycosylated haemoglobin (glyhb) and dropped glyhb itself, along with the 13 rows where it was missing.",
                   "Height, waist and hip were converted from inches to centimetres; weight from pounds to kilograms.",
                   "Dropped the mostly empty second blood-pressure reading, the subject ID, the cholesterol/HDL ratio (computable from two other columns) and the time of the blood draw.",
                   "Original study: Willems JP et al. (1997) Prevalence of coronary heart disease risk factors among rural blacks. Journal of Community Health 22:365-375."],
        },
    })


# ── Titanic (titanic3) ───────────────────────────────────────────────────
def titanic():
    h, rows = read(os.path.join(RAW, "titanic3.csv"))
    i = {c: h.index(c) for c in h}
    keep = ["pclass", "sex", "age", "sibsp", "parch", "fare", "embarked", "survived"]
    write("titanic.csv", keep, [[r[i[c]] for c in keep] for r in rows])
    recipe({
        "id": "titanic",
        "csv_name": "titanic.csv",
        "source_url": "https://hbiostat.org/data/",
        "license": "Free to use with attribution - \"Data obtained from http://hbiostat.org/data courtesy of the Vanderbilt University Department of Biostatistics.\"",
        "title": {"ko": "타이타닉 생존자", "en": "Titanic Survivors"},
        "target": {"ko": "생존 여부 (survived)", "en": "survived"},
        "headers": [
            H("pclass", "객실 등급 (pclass)"),
            H("sex", "성별 (sex)"),
            H("age", "나이 (age)"),
            H("sibsp", "함께 탄 형제·배우자 수 (sibsp)"),
            H("parch", "함께 탄 부모·자녀 수 (parch)"),
            H("fare", "운임 (fare, 파운드)"),
            H("embarked", "탑승 항구 (embarked)"),
            H("survived", "생존 여부 (survived)"),
        ],
        "values": {
            "pclass": {"ko": {"1": "1등석 (1)", "2": "2등석 (2)", "3": "3등석 (3)"},
                       "en": {"1": "1st class (1)", "2": "2nd class (2)", "3": "3rd class (3)"}},
            "sex": {"ko": {"male": "남성 (male)", "female": "여성 (female)"}},
            "embarked": {"ko": {"S": "사우샘프턴 (S)", "C": "셰르부르 (C)", "Q": "퀸스타운 (Q)", "": ""},
                         "en": {"S": "Southampton (S)", "C": "Cherbourg (C)", "Q": "Queenstown (Q)", "": ""}},
            "survived": {"ko": {"0": "사망 (0)", "1": "생존 (1)"},
                         "en": {"0": "Died (0)", "1": "Survived (1)"}},
        },
        "body": {
            "ko": "1912년 타이타닉호 승객 1,309명의 기록입니다. 객실 등급·성별·나이·동승 가족·운임으로 생존 여부를 맞힙니다.\n"
                  "\"여성과 아이 먼저\"가 실제로 지켜졌는지, 객실 등급이 생사를 갈랐는지 모델에게 물어볼 수 있습니다. 나이가 비어 있는 승객이 약 260명 있습니다.",
            "en": "Records of 1,309 passengers on the Titanic in 1912. Predict survival from class, sex, age, family aboard and fare.\n"
                  "Ask the model whether 'women and children first' really held, and whether class decided who lived. About 260 passengers have no recorded age.",
        },
        "notes": {
            "ko": ["구명보트 번호(boat)와 시신 번호(body)를 뺐습니다. 생존 여부를 그대로 알려 주는 열입니다.",
                   "이름, 표 번호, 객실 번호, 목적지는 학습에 쓰기 어려운 글자 열이라 뺐습니다.",
                   "생존 여부를 맨 끝 열로 옮겼습니다.",
                   "원 자료: Thomas Cason (UVa)이 Encyclopedia Titanica를 바탕으로 정리한 titanic3, Frank E. Harrell Jr. 배포."],
            "en": ["Dropped the lifeboat number (boat) and body number (body); both give away whether a passenger survived.",
                   "Dropped name, ticket, cabin and destination - free-text columns that a model cannot use directly.",
                   "Moved survived to the last column.",
                   "Original: titanic3, compiled by Thomas Cason (UVa) from Encyclopedia Titanica and distributed by Frank E. Harrell Jr."],
        },
    })


# ── Palmer Penguins ──────────────────────────────────────────────────────
def penguins():
    h, rows = read(os.path.join(RAW, "penguins.csv"))
    i = {c: h.index(c) for c in h}
    order = ["island", "bill_length_mm", "bill_depth_mm", "flipper_length_mm", "body_mass_g", "sex", "year", "species"]
    write("penguins.csv", order, [["" if r[i[c]] == "NA" else r[i[c]] for c in order] for r in rows])
    recipe({
        "id": "penguins",
        "csv_name": "penguins.csv",
        "source_url": "https://allisonhorst.github.io/palmerpenguins/",
        "license": "CC0 1.0",
        "title": {"ko": "남극 펭귄", "en": "Palmer Penguins"},
        "target": {"ko": "펭귄 종 (species)", "en": "species"},
        "headers": [
            H("island", "서식 섬 (island)"),
            H("bill_length_mm", "부리 길이 (bill_length_mm)"),
            H("bill_depth_mm", "부리 깊이 (bill_depth_mm)"),
            H("flipper_length_mm", "날개 길이 (flipper_length_mm)"),
            H("body_mass_g", "몸무게 (body_mass_g)"),
            H("sex", "성별 (sex)"),
            H("year", "관측 연도 (year)"),
            H("species", "펭귄 종 (species)"),
        ],
        "values": {
            "island": {"ko": {"Biscoe": "비스코 섬 (Biscoe)", "Dream": "드림 섬 (Dream)", "Torgersen": "토르거센 섬 (Torgersen)"}},
            "sex": {"ko": {"male": "수컷 (male)", "female": "암컷 (female)", "": ""}},
            "species": {"ko": {"Adelie": "아델리 (Adelie)", "Chinstrap": "턱끈 (Chinstrap)", "Gentoo": "젠투 (Gentoo)"}},
        },
        "body": {
            "ko": "남극 파머 기지 근처 세 섬에서 관측한 펭귄 344마리의 기록입니다. 부리·날개 길이와 몸무게로 세 종을 구별합니다.\n"
                  "몸무게를 목표로 하면 회귀, 종 열을 빼면 군집으로도 쓸 수 있습니다. 붓꽃 데이터의 다음 단계로 알맞습니다.",
            "en": "344 penguins observed on three islands near Palmer Station, Antarctica. Tell three species apart by bill, flipper and body mass.\n"
                  "Use body mass as the target for regression, or drop species for clustering - a natural next step after Iris.",
        },
        "notes": {
            "ko": ["원본의 결측 표시 NA를 빈칸으로 바꾸었습니다. 측정값이 모두 빈 2행과 성별만 빈 9행이 있습니다.",
                   "펭귄 종을 맨 끝 열로 옮겼습니다.",
                   "원 자료: Gorman KB, Williams TD, Fraser WR (2014). PLoS ONE 9(3):e90081. Palmer Station Antarctica LTER."],
            "en": ["Replaced the NA marker with empty cells. Two rows have no measurements at all; nine more are missing only sex.",
                   "Moved species to the last column.",
                   "Original: Gorman KB, Williams TD, Fraser WR (2014). PLoS ONE 9(3):e90081. Palmer Station Antarctica LTER."],
        },
    })


# ── Zoo ──────────────────────────────────────────────────────────────────
def zoo():
    _, rows = read(os.path.join(RAW, "zoo", "zoo.data"), header=False)
    zh = ["animal_name", "hair", "feathers", "eggs", "milk", "airborne", "aquatic", "predator",
          "toothed", "backbone", "breathes", "venomous", "fins", "legs", "tail", "domestic", "catsize", "type"]
    write("zoo.csv", zh, rows)
    zko = {"hair": "털", "feathers": "깃털", "eggs": "알을 낳음", "milk": "젖을 먹임", "airborne": "날 수 있음",
           "aquatic": "물에 삶", "predator": "포식자", "toothed": "이빨", "backbone": "척추", "breathes": "공기 호흡",
           "venomous": "독", "fins": "지느러미", "tail": "꼬리", "domestic": "가축·반려동물", "catsize": "고양이보다 큼"}
    yn = {"ko": {"0": "아니오 (0)", "1": "예 (1)"}, "en": {"0": "No (0)", "1": "Yes (1)"}}
    zhead = [H("animal_name", "동물 이름 (animal_name)")]
    for c in zh[1:-1]:
        zhead.append(H(c, "다리 수 (legs)" if c == "legs" else f"{zko[c]} ({c})"))
    zhead.append(H("type", "동물 분류 (type)"))
    recipe({
        "id": "zoo",
        "csv_name": "zoo.csv",
        "source_url": "https://archive.ics.uci.edu/dataset/111/zoo",
        "license": "CC BY 4.0",
        "title": {"ko": "동물원", "en": "Zoo"},
        "target": {"ko": "동물 분류 (type)", "en": "type"},
        "headers": zhead,
        "values": dict({c: yn for c in zko}, type={
            "ko": {"1": "포유류 (1)", "2": "조류 (2)", "3": "파충류 (3)", "4": "어류 (4)",
                   "5": "양서류 (5)", "6": "곤충 (6)", "7": "무척추동물 (7)"},
            "en": {"1": "Mammal (1)", "2": "Bird (2)", "3": "Reptile (3)", "4": "Fish (4)",
                   "5": "Amphibian (5)", "6": "Insect (6)", "7": "Invertebrate (7)"}}),
        "body": {
            "ko": "동물 101종을 털·깃털·알·젖·다리 수 같은 특징 열여섯 가지로 적은 표입니다. 포유류·조류·어류 등 일곱 무리로 분류합니다.\n"
                  "결정 트리를 그려 보면 '젖을 먹이면 포유류'처럼 생물 시간에 배운 기준이 그대로 나타납니다. 오리너구리나 돌고래가 어디로 가는지 확인해 보세요.",
            "en": "101 animals described by sixteen traits - hair, feathers, eggs, milk, number of legs and so on - sorted into seven groups such as mammals, birds and fish.\n"
                  "A decision tree rediscovers biology-class rules like 'gives milk -> mammal'. See where the platypus and the dolphin end up.",
        },
        "notes": {
            "ko": ["원본에는 헤더가 없어 문서(zoo.names)의 순서대로 열 이름을 붙였습니다.",
                   "0/1로 적힌 특징은 예/아니오로, 분류 번호 1~7은 무리 이름으로 풀어 적었습니다.",
                   "동물 이름 열은 학습에 쓰지 말고 결과를 확인할 때만 보세요. 원본 그대로 frog가 두 번(독이 있는 개구리와 없는 개구리), girl이 포유류로 들어 있습니다."],
            "en": ["The original has no header; column names follow the order in zoo.names.",
                   "0/1 traits are spelled out as Yes/No, and the class numbers 1-7 as group names.",
                   "Leave animal_name out of training and use it only to check results. As in the original, 'frog' appears twice (venomous and not) and 'girl' is listed as a mammal."],
        },
    })


# ── Mushroom ─────────────────────────────────────────────────────────────
def mushroom():
    _, rows = read(os.path.join(RAW, "mushroom", "agaricus-lepiota.data"), header=False)
    COLOR = {"k": ("black", "검정"), "n": ("brown", "갈색"), "b": ("buff", "담황색"), "h": ("chocolate", "초콜릿색"),
             "g": ("gray", "회색"), "r": ("green", "초록"), "o": ("orange", "주황"), "p": ("pink", "분홍"),
             "u": ("purple", "보라"), "e": ("red", "빨강"), "w": ("white", "흰색"), "y": ("yellow", "노랑"),
             "c": ("cinnamon", "계피색")}
    SURF = {"f": ("fibrous", "섬유질"), "y": ("scaly", "비늘"), "k": ("silky", "비단결"), "s": ("smooth", "매끈함")}


    def pick(d, keys):
        return {k: d[k] for k in keys}


    M = [  # (en column, ko name, {code: (en, ko)})
        ("cap-shape", "갓 모양", {"b": ("bell", "종 모양"), "c": ("conical", "원뿔형"), "x": ("convex", "볼록형"),
                                "f": ("flat", "평평함"), "k": ("knobbed", "혹 있음"), "s": ("sunken", "움푹함")}),
        ("cap-surface", "갓 표면", {"f": ("fibrous", "섬유질"), "g": ("grooves", "홈 있음"), "y": ("scaly", "비늘"), "s": ("smooth", "매끈함")}),
        ("cap-color", "갓 색", pick(COLOR, "nbcgrpuewy")),
        ("bruises", "멍 자국", {"t": ("bruises", "있음"), "f": ("no", "없음")}),
        ("odor", "냄새", {"a": ("almond", "아몬드"), "l": ("anise", "아니스"), "c": ("creosote", "크레오소트"),
                         "y": ("fishy", "비린내"), "f": ("foul", "악취"), "m": ("musty", "곰팡내"), "n": ("none", "없음"),
                         "p": ("pungent", "톡 쏘는 냄새"), "s": ("spicy", "매운 냄새")}),
        ("gill-attachment", "주름살 붙음새", {"a": ("attached", "붙음"), "d": ("descending", "내려붙음"),
                                           "f": ("free", "떨어짐"), "n": ("notched", "홈파진")}),
        ("gill-spacing", "주름살 간격", {"c": ("close", "촘촘함"), "w": ("crowded", "빽빽함"), "d": ("distant", "성김")}),
        ("gill-size", "주름살 폭", {"b": ("broad", "넓음"), "n": ("narrow", "좁음")}),
        ("gill-color", "주름살 색", pick(COLOR, "knbhgropuewy")),
        ("stalk-shape", "자루 모양", {"e": ("enlarging", "아래로 굵어짐"), "t": ("tapering", "아래로 가늘어짐")}),
        ("stalk-root", "자루 뿌리", {"b": ("bulbous", "구근형"), "c": ("club", "곤봉형"), "u": ("cup", "컵형"),
                                   "e": ("equal", "굵기 같음"), "z": ("rhizomorphs", "뿌리꼴균사"), "r": ("rooted", "뿌리형"),
                                   "?": ("", "")}),
        ("stalk-surface-above-ring", "턱받이 위 자루 표면", SURF),
        ("stalk-surface-below-ring", "턱받이 아래 자루 표면", SURF),
        ("stalk-color-above-ring", "턱받이 위 자루 색", pick(COLOR, "nbcgopewy")),
        ("stalk-color-below-ring", "턱받이 아래 자루 색", pick(COLOR, "nbcgopewy")),
        ("veil-type", "내피막 종류", {"p": ("partial", "부분"), "u": ("universal", "전체")}),
        ("veil-color", "내피막 색", pick(COLOR, "nowy")),
        ("ring-number", "턱받이 수", {"n": ("none", "없음"), "o": ("one", "하나"), "t": ("two", "둘")}),
        ("ring-type", "턱받이 모양", {"c": ("cobwebby", "거미줄형"), "e": ("evanescent", "사라짐"), "f": ("flaring", "벌어짐"),
                                     "l": ("large", "큼"), "n": ("none", "없음"), "p": ("pendant", "늘어짐"),
                                     "s": ("sheathing", "칼집형"), "z": ("zone", "띠형")}),
        ("spore-print-color", "포자 색", pick(COLOR, "knbhrouwy")),
        ("population", "자라는 모습", {"a": ("abundant", "아주 많음"), "c": ("clustered", "모여 남"), "n": ("numerous", "많음"),
                                      "s": ("scattered", "흩어짐"), "v": ("several", "여럿"), "y": ("solitary", "홀로")}),
        ("habitat", "서식지", {"g": ("grasses", "풀밭"), "l": ("leaves", "낙엽"), "m": ("meadows", "목초지"),
                              "p": ("paths", "길가"), "u": ("urban", "도시"), "w": ("waste", "황무지"), "d": ("woods", "숲")}),
        ("class", "식용 여부", {"e": ("edible", "식용"), "p": ("poisonous", "독버섯")}),
    ]
    write("mushroom.csv", [c for c, _, _ in M], [r[1:] + r[:1] for r in rows])


    def vm(codes):
        ko = {k: (f"{v[1]} ({v[0]})" if v[0] else "") for k, v in codes.items()}
        en = {k: v[0] for k, v in codes.items()}
        return {"ko": ko, "en": en}


    recipe({
        "id": "mushroom",
        "csv_name": "mushroom.csv",
        "source_url": "https://archive.ics.uci.edu/dataset/73/mushroom",
        "license": "CC BY 4.0",
        "title": {"ko": "독버섯 판별", "en": "Mushroom"},
        "target": {"ko": "식용 여부 (class)", "en": "class"},
        "headers": [H(c, f"{k} ({c})") for c, k, _ in M],
        "values": {c: vm(codes) for c, _, codes in M},
        "body": {
            "ko": "북미 버섯 도감에 실린 주름버섯류 23종을 갓·주름살·자루·냄새 등 스물두 가지 특징으로 적은 8,124개 표본입니다. 먹을 수 있는지, 독버섯인지 맞힙니다.\n"
                  "모든 열이 숫자가 아닌 범주라 결정 트리가 특히 잘 맞습니다. 냄새 하나만으로 얼마나 맞힐 수 있는지 확인해 보세요. 실제 버섯을 이 모델로 판단해서는 안 됩니다.",
            "en": "8,124 samples of 23 gilled-mushroom species from a North American field guide, described by 22 traits of cap, gills, stalk and odour. Predict edible or poisonous.\n"
                  "Every column is categorical, which suits decision trees. Try predicting from odour alone. Never use this model on a real mushroom.",
        },
        "notes": {
            "ko": ["원본에는 헤더가 없고 값이 모두 알파벳 한 글자 코드입니다. 문서(agaricus-lepiota.names)를 따라 열 이름을 붙이고 코드를 낱말로 풀어 적었습니다.",
                   "자루 뿌리(stalk-root)의 결측 표시 ?를 빈칸으로 바꾸었습니다 (2,480행).",
                   "식용 여부를 맨 앞에서 맨 끝 열로 옮겼습니다. 원본에서 '먹을 수 있는지 알 수 없음'은 독버섯으로 분류되어 있습니다.",
                   "내피막 종류(veil-type)는 모든 행이 partial로 같습니다. 원본 그대로 두었습니다."],
            "en": ["The original has no header and every value is a one-letter code. Column names and codes were spelled out following agaricus-lepiota.names.",
                   "The missing-value marker ? in stalk-root was replaced with an empty cell (2,480 rows).",
                   "Moved class from the first column to the last. In the original, 'edibility unknown' was folded into poisonous.",
                   "veil-type is 'partial' in every row; it is kept as in the original."],
        },
    })


# ── Combined Cycle Power Plant ───────────────────────────────────────────
def power_plant():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(RAW, "ccpp", "CCPP", "Folds5x2_pp.xlsx"), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))  # 시트 다섯 장은 같은 자료를 섞은 순서만 다르다
    assert list(rows[0]) == ["AT", "V", "AP", "RH", "PE"], rows[0]
    write("power-plant.csv", ["AT", "V", "AP", "RH", "PE"], [[str(c) for c in r] for r in rows[1:]])
    recipe({
        "id": "power-plant",
        "csv_name": "power-plant.csv",
        "source_url": "https://archive.ics.uci.edu/dataset/294/combined+cycle+power+plant",
        "license": "CC BY 4.0",
        "title": {"ko": "발전소 출력", "en": "Combined Cycle Power Plant"},
        "target": {"ko": "발전량 (PE, MW)", "en": "PE"},
        "headers": [
            H("AT", "기온 (AT, °C)"),
            H("V", "배기 진공도 (V, cmHg)"),
            H("AP", "기압 (AP, mbar)"),
            H("RH", "상대습도 (RH, %)"),
            H("PE", "발전량 (PE, MW)"),
        ],
        "body": {
            "ko": "터키의 한 복합화력발전소가 6년 동안 전부하로 돌 때 한 시간 평균으로 기록한 9,568개 값입니다. 바깥 기온·기압·습도와 배기 진공도로 그 시간의 발전량을 예측합니다.\n"
                  "공기가 더우면 밀도가 낮아져 가스터빈이 들이마시는 공기가 줄고, 발전량도 줄어듭니다. 기온 하나만으로도 꽤 정확한 직선이 그려집니다.\n\n"
                  "- AT: 바깥 기온 (°C)\n- V: 증기터빈 배기 쪽 진공도 (cmHg). 기온이 오르면 함께 오르는 경향이 있습니다.\n"
                  "- AP: 대기압 (mbar, = hPa)\n- RH: 상대습도 (%)\n- PE: 순 발전량 (MW)",
            "en": "9,568 hourly averages from a combined-cycle power plant in Turkey running at full load over six years. Predict the hour's power output from outdoor temperature, pressure, humidity and exhaust vacuum.\n"
                  "Hot air is thin, so the gas turbine draws in less of it and output drops. Temperature alone already gives a fairly tight straight line.\n\n"
                  "- AT: ambient temperature (°C)\n- V: exhaust vacuum on the steam-turbine side (cmHg); it tends to rise with temperature\n"
                  "- AP: ambient pressure (mbar, = hPa)\n- RH: relative humidity (%)\n- PE: net electrical output (MW)",
        },
        "notes": {
            "ko": ["원본 엑셀의 첫 번째 시트를 CSV로 옮겼습니다. 다섯 시트는 같은 자료의 행 순서만 다릅니다.",
                   "원 연구: Tüfekci P. (2014) Prediction of full load electrical power output of a base load operated combined cycle power plant using machine learning methods. International Journal of Electrical Power & Energy Systems 60:126-140."],
            "en": ["Converted the first sheet of the original workbook to CSV; the five sheets hold the same rows in different orders.",
                   "Original study: Tüfekci P. (2014) Prediction of full load electrical power output of a base load operated combined cycle power plant using machine learning methods. International Journal of Electrical Power & Energy Systems 60:126-140."],
        },
    })


# ── Real Estate Valuation (Taiwan) ───────────────────────────────────────
def real_estate():
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(RAW, "realestate", "Real estate valuation data set.xlsx"), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    assert rows[0][0] == "No" and rows[0][-1].startswith("Y"), rows[0]

    def ym(x):  # 2012.917 -> 2012년 11월. 원본은 (월 - 1) / 12 가 아니라 월 / 12 로 적었다.
        y = int(x)
        m = round((x - y) * 12)
        return (y, m) if m else (y - 1, 12)

    out = []
    for r in rows[1:]:
        y, m = ym(r[1])
        out.append([f"{y}-{m:02d}", f"{r[2]:g}", f"{r[3]:.1f}", str(r[4]), f"{r[5]:.5f}", f"{r[6]:.5f}", f"{r[7]:g}"])
    write("real-estate.csv",
          ["transaction month", "house age (years)", "distance to MRT (m)", "convenience stores",
           "latitude", "longitude", "price per ping (10k NTD)"], out)
    recipe({
        "id": "real-estate",
        "csv_name": "real-estate.csv",
        "source_url": "https://archive.ics.uci.edu/dataset/477/real+estate+valuation+data+set",
        "license": "CC BY 4.0",
        "title": {"ko": "대만 집값", "en": "Taiwan Real Estate Valuation"},
        "target": {"ko": "평당 가격 (price per ping, 만 대만달러)", "en": "price per ping (10k NTD)"},
        "headers": [
            H("transaction month", "거래 연월 (transaction month)"),
            H("house age (years)", "집 나이 (house age, 년)"),
            H("distance to MRT (m)", "가장 가까운 지하철역까지 거리 (distance to MRT, m)"),
            H("convenience stores", "걸어갈 수 있는 편의점 수 (convenience stores)"),
            H("latitude", "위도 (latitude)"),
            H("longitude", "경도 (longitude)"),
            H("price per ping (10k NTD)", "평당 가격 (price per ping, 만 대만달러)"),
        ],
        "body": {
            "ko": "대만 신베이시 신뎬구에서 2012~2013년에 거래된 집 414채의 기록입니다. 집 나이, 지하철역까지 거리, 주변 편의점 수로 집값을 예측합니다.\n"
                  "\"역이 가깝고 편의점이 많으면 비싸다, 오래된 집은 싸다\"는 짐작이 맞는지 확인해 보세요.\n\n"
                  "가격 단위는 1평(약 3.3㎡)당 만 대만달러입니다. 예를 들어 37.9는 평당 37만 9천 대만달러입니다.\n"
                  "거래 연월은 하나하나가 다른 값이 아니라 12개월 중 하나이므로 범주로 써도 됩니다. 위도·경도는 지도 위치이므로 군집에 써 볼 만합니다.",
            "en": "414 home sales in Xindian District, New Taipei City, Taiwan, in 2012-2013. Predict price from house age, distance to the nearest MRT station and the number of convenience stores nearby.\n"
                  "Check the hunch that near-the-station, many-shops homes cost more and old homes cost less.\n\n"
                  "Price is in units of 10,000 New Taiwan dollars per ping (about 3.3 m²): 37.9 means 379,000 NTD per ping.\n"
                  "The transaction month takes only twelve values, so it can be treated as a category. Latitude and longitude place each home on the map and are worth trying for clustering.",
        },
        "notes": {
            "ko": ["원본의 일련번호 열(No)을 뺐습니다.",
                   "거래 시점이 2012.917처럼 연도의 소수로 적혀 있어 2012-11 같은 연월로 바꾸었습니다 (소수 부분 × 12 = 월).",
                   "원본 열 이름의 X1~X6, Y 기호를 떼고 단위를 열 이름에 적었습니다.",
                   "원 연구: Yeh I-C, Hsu T-K (2018) Building real estate valuation models with comparative approach through case-based reasoning. Applied Soft Computing 65:260-271."],
            "en": ["Dropped the row-number column (No).",
                   "The transaction date was a decimal year such as 2012.917; it is written as a year-month such as 2012-11 (fraction × 12 = month).",
                   "Removed the X1-X6/Y prefixes from the column names and put units in the names.",
                   "Original study: Yeh I-C, Hsu T-K (2018) Building real estate valuation models with comparative approach through case-based reasoning. Applied Soft Computing 65:260-271."],
        },
    })


# ── Beijing PM2.5 ────────────────────────────────────────────────────────
PM_LEVELS = [(15, "good"), (35, "moderate"), (75, "bad")]  # 한국 환경부 초미세먼지 예보 등급


def beijing_pm25():
    h, rows = read(os.path.join(RAW, "pm25", "PRSA_data_2010.1.1-2014.12.31.csv"))
    i = {c: h.index(c) for c in h}
    out, dropped = [], 0
    for r in rows:
        pm = r[i["pm2.5"]]
        if pm == "NA":
            dropped += 1
            continue
        v = float(pm)
        level = next((name for top, name in PM_LEVELS if v <= top), "very bad")
        out.append([r[i[c]] for c in ("year", "month", "day", "hour", "TEMP", "DEWP", "PRES", "cbwd", "Iws", "Is", "Ir")]
                   + [pm, level])
    assert dropped == 2067, dropped
    write("beijing-pm25.csv",
          ["year", "month", "day", "hour", "TEMP", "DEWP", "PRES", "cbwd", "Iws", "Is", "Ir", "pm2.5", "pm2.5 level"], out)
    level_ko = {"good": "좋음 (good)", "moderate": "보통 (moderate)", "bad": "나쁨 (bad)", "very bad": "매우 나쁨 (very bad)"}
    recipe({
        "id": "beijing-pm25",
        "csv_name": "beijing-pm25.csv",
        "source_url": "https://archive.ics.uci.edu/dataset/381/beijing+pm2+5+data",
        "license": "CC BY 4.0",
        "title": {"ko": "베이징 초미세먼지", "en": "Beijing PM2.5"},
        "target": {"ko": "초미세먼지 농도 (pm2.5) 또는 등급 (pm2.5 level)", "en": "pm2.5 or pm2.5 level"},
        "headers": [
            H("year", "연 (year)"), H("month", "월 (month)"), H("day", "일 (day)"), H("hour", "시 (hour)"),
            H("TEMP", "기온 (TEMP, °C)"),
            H("DEWP", "이슬점 (DEWP, °C)"),
            H("PRES", "기압 (PRES, hPa)"),
            H("cbwd", "풍향 (cbwd)"),
            H("Iws", "누적 풍속 (Iws, m/s)"),
            H("Is", "눈 누적 시간 (Is, 시간)"),
            H("Ir", "비 누적 시간 (Ir, 시간)"),
            H("pm2.5", "초미세먼지 농도 (pm2.5, μg/m³)"),
            H("pm2.5 level", "초미세먼지 등급 (pm2.5 level)"),
        ],
        "values": {
            "cbwd": {"ko": {"NE": "북동풍 (NE)", "NW": "북서풍 (NW)", "SE": "남동풍 (SE)", "cv": "바람 약하고 바뀜 (cv)"},
                     "en": {"NE": "northeast (NE)", "NW": "northwest (NW)", "SE": "southeast (SE)", "cv": "calm and variable (cv)"}},
            "pm2.5 level": {"ko": level_ko},
        },
        "body": {
            "ko": "2010~2014년 베이징 주재 미국 대사관에서 잰 시간별 초미세먼지(PM2.5) 농도와 베이징 공항의 기상 관측을 묶은 자료입니다. 기온·이슬점·기압·바람·눈비로 초미세먼지를 예측합니다.\n"
                  "겨울에, 밤에, 바람이 잦아들 때 나빠진다는 짐작을 확인해 보세요. 북서풍이 불면 깨끗해지는 것도 보입니다.\n\n"
                  "## 목표 열이 둘입니다 - 하나를 고르고 나머지는 빼십시오\n\n"
                  "- `초미세먼지 농도 (pm2.5, μg/m³)` - 회귀\n"
                  "- `초미세먼지 등급 (pm2.5 level)` - 분류. 한국 환경부 예보 기준으로 15 이하 좋음, 35 이하 보통, 75 이하 나쁨, 그 위는 매우 나쁨입니다.\n\n"
                  "둘 다 넣으면 농도로 등급을 맞히는 셈이라, 모델은 아무것도 배우지 않고 100%에 가까운 점수를 냅니다.\n\n"
                  "## 헷갈리는 열\n\n"
                  "- 누적 풍속: 같은 풍향이 이어진 동안의 풍속을 계속 더한 값입니다. 바람이 오래 같은 방향으로 불수록 커집니다.\n"
                  "- 눈·비 누적 시간: 눈이나 비가 몇 시간째 이어지고 있는지입니다.\n"
                  "- 이슬점: 공기 속 수증기가 많을수록 높습니다. 기온과 가까우면 습한 날입니다.",
            "en": "Hourly PM2.5 measured at the US Embassy in Beijing, 2010-2014, joined with weather observed at Beijing airport. Predict fine dust from temperature, dew point, pressure, wind and snow or rain.\n"
                  "Check the hunch that it gets worse in winter, at night and when the wind dies down - and that a northwest wind clears it.\n\n"
                  "## There are two target columns - pick one and drop the other\n\n"
                  "- `pm2.5` (μg/m³) - regression\n"
                  "- `pm2.5 level` - classification, using the Korean Ministry of Environment forecast bands: up to 15 good, up to 35 moderate, up to 75 bad, above that very bad.\n\n"
                  "Keeping both means predicting the band from the number, and the model scores near 100% having learned nothing.\n\n"
                  "## Columns that need explaining\n\n"
                  "- Iws: cumulated wind speed - the wind speed summed over the hours the same direction has persisted.\n"
                  "- Is, Ir: how many hours snow or rain has been falling so far.\n"
                  "- DEWP: dew point; the closer it is to the temperature, the more humid the air.",
        },
        "notes": {
            "ko": ["초미세먼지 값이 비어 있는(NA) 2,067시간을 뺐습니다. 목표가 없는 행은 학습에 쓸 수 없습니다.",
                   "원본의 일련번호 열(No)을 뺐습니다.",
                   "한국 환경부 초미세먼지 예보 기준으로 등급 열을 새로 만들었습니다.",
                   "목표 열을 맨 끝으로 옮겼습니다.",
                   "원 연구: Liang X. et al. (2015) Assessing Beijing's PM2.5 pollution: severity, weather impact, APEC and winter heating. Proceedings of the Royal Society A 471:20150257."],
            "en": ["Dropped the 2,067 hours with no PM2.5 reading (NA); rows without a target cannot be used for training.",
                   "Dropped the row-number column (No).",
                   "Added a level column using the Korean Ministry of Environment PM2.5 forecast bands.",
                   "Moved the targets to the end.",
                   "Original study: Liang X. et al. (2015) Assessing Beijing's PM2.5 pollution: severity, weather impact, APEC and winter heating. Proceedings of the Royal Society A 471:20150257."],
        },
    })


# ── 전국 전력수요 + 기온 ─────────────────────────────────────────────────
# 전력거래소 자료에는 기온이 없다. 다섯 대도시의 시간별 기온을 Open-Meteo에서 받아
# 평균을 낸 뒤 같은 시각에 붙인다. 한 도시만 쓰면 그 도시 날씨만 반영된다.
DEMAND_YEAR = 2025
DEMAND_CITIES = [("서울", 37.5665, 126.9780), ("부산", 35.1796, 129.0756), ("대구", 35.8714, 128.6014),
                 ("광주", 35.1595, 126.8526), ("대전", 36.3504, 127.3845)]
# 2025년 대한민국 공휴일 (대체공휴일·임시공휴일 포함). 해가 바뀌면 고친다.
DEMAND_HOLIDAYS = {
    "2025-01-01", "2025-01-27", "2025-01-28", "2025-01-29", "2025-01-30",
    "2025-03-01", "2025-03-03", "2025-05-05", "2025-05-06", "2025-06-03", "2025-06-06",
    "2025-08-15", "2025-10-03", "2025-10-05", "2025-10-06", "2025-10-07", "2025-10-08",
    "2025-10-09", "2025-12-25",
}
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
WEEKDAYS_KO = ["월", "화", "수", "목", "금", "토", "일"]


def power_demand():
    import datetime as dt, urllib.parse, urllib.request, time
    with open(os.path.join(RAW, "kpx.csv"), encoding="cp949", newline="") as f:
        kpx = [r for r in csv.reader(f) if any(c.strip() for c in r)]   # 원본 끝에 쉼표만 있는 줄이 있다
    assert kpx[0][0] == "날짜" and len(kpx[0]) == 25, kpx[0]

    # 시각마다 다섯 도시 기온의 평균. 1시~24시를 덮으려면 다음 해 1월 1일 0시까지 필요하다.
    temps = {}
    for name, lat, lon in DEMAND_CITIES:
        q = urllib.parse.urlencode({"latitude": lat, "longitude": lon, "hourly": "temperature_2m",
                                    "start_date": f"{DEMAND_YEAR}-01-01", "end_date": f"{DEMAND_YEAR + 1}-01-01",
                                    "timezone": "Asia/Seoul"})
        with urllib.request.urlopen("https://archive-api.open-meteo.com/v1/archive?" + q, timeout=120) as r:
            hourly = json.load(r)["hourly"]
        for t, v in zip(hourly["time"], hourly["temperature_2m"]):
            temps.setdefault(t, []).append(v)
        time.sleep(1)

    out = []
    for r in kpx[1:]:
        d = dt.date.fromisoformat(r[0])
        for hour in range(1, 25):
            # '1시'는 0시~1시 사이의 수요다. 그 구간이 끝나는 시각(1시, 24시는 다음 날 0시)의 기온을 붙인다.
            stamp = (dt.datetime.combine(d, dt.time()) + dt.timedelta(hours=hour)).strftime("%Y-%m-%dT%H:%M")
            vals = temps[stamp]
            assert len(vals) == len(DEMAND_CITIES) and None not in vals, stamp
            out.append([r[0], str(d.month), WEEKDAYS[d.weekday()],
                        "yes" if r[0] in DEMAND_HOLIDAYS else "no",
                        str(hour), f"{sum(vals) / len(vals):.1f}", r[hour].replace(",", "").strip()])
    assert len(out) == 365 * 24, len(out)
    write("power-demand.csv", ["date", "month", "weekday", "holiday", "hour", "temperature (C)", "demand (MWh)"], out)

    city_list_ko = "·".join(n for n, _, _ in DEMAND_CITIES)
    recipe({
        "id": "power-demand",
        "csv_name": "power-demand.csv",
        "source_url": {"ko": "https://www.data.go.kr/data/15065266/fileData.do (전력수요, 한국전력거래소) + https://open-meteo.com/ (기온)",
                       "en": "https://www.data.go.kr/data/15065266/fileData.do (demand, Korea Power Exchange) + https://open-meteo.com/ (temperature)"},
        "license": {"ko": "전력수요: 공공데이터포털 이용허락범위 제한 없음. 기온: CC BY 4.0 (Open-Meteo)",
                    "en": "Demand: no usage restrictions per data.go.kr. Temperature: CC BY 4.0 (Open-Meteo)"},
        "title": {"ko": "전국 전력수요", "en": "Korea Hourly Power Demand"},
        "target": {"ko": "전력수요 (demand, MWh)", "en": "demand (MWh)"},
        "headers": [
            H("date", "날짜 (date)"),
            H("month", "월 (month)"),
            H("weekday", "요일 (weekday)"),
            H("holiday", "공휴일 (holiday)"),
            H("hour", "시각 (hour)"),
            H("temperature (C)", "기온 (temperature, °C)"),
            H("demand (MWh)", "전력수요 (demand, MWh)"),
        ],
        "values": {
            "weekday": {"ko": {e: f"{k} ({e})" for e, k in zip(WEEKDAYS, WEEKDAYS_KO)}},
            "holiday": {"ko": {"yes": "공휴일 (yes)", "no": "평일·주말 (no)"}},
        },
        "body": {
            "ko": f"{DEMAND_YEAR}년 한 해 동안 한 시간마다 전국(육지와 제주)에서 쓴 전력량과 그 시각의 기온입니다. 시각·요일·공휴일·기온으로 전력수요를 예측합니다.\n"
                  "낮에 많고 새벽에 적다, 주말과 공휴일에는 줄어든다, 한여름과 한겨울에는 냉난방 때문에 늘어난다 - 학습 전에 짐작해 보고 모델과 비교해 보세요.\n"
                  "기온과 전력수요로 산점도를 그리면 봄·가을에 낮고 양쪽으로 올라가는 U자가 보입니다. 직선 하나로는 잘 맞지 않는 이유를 생각해 볼 수 있습니다.\n\n"
                  "`날짜 (date)`는 하루에 24행씩 같은 값이므로 학습에서 빼고, 학습·테스트를 시간 순서로 나눌 때 쓰십시오.\n\n"
                  "## 이렇게 만들었습니다 - 두 자료를 합쳤습니다\n\n"
                  f"1. **전력수요**: 한국전력거래소가 공공데이터포털에 올린 \"시간별 전국 전력수요량\" {DEMAND_YEAR}년 파일입니다. 원본은 하루가 한 행이고 1시~24시가 열로 늘어선 표(365행 × 25열)라서, 하루 24행씩 세로로 펼쳤습니다. 원본 설명에 따르면 발전단 기준이며 수요예측용 잠정자료입니다.\n"
                  f"2. **기온**: 전력거래소 자료에는 기온이 없습니다. Open-Meteo 과거 기상 API에서 {city_list_ko} 다섯 도시의 시간별 기온(지상 2m)을 받아 같은 시각끼리 평균을 냈습니다. 한 도시만 쓰면 그 도시의 날씨만 반영되기 때문입니다.\n"
                  "3. **시각 맞추기**: 원본의 '1시'는 0시부터 1시까지 쓴 양입니다. 그 구간이 끝나는 시각의 기온을 붙였습니다. '24시'에는 다음 날 0시의 기온이 붙습니다.\n"
                  "4. **달력 열**: 날짜에서 월과 요일을 계산했고, 공휴일은 대체공휴일과 임시공휴일(1월 27일, 6월 3일 대통령 선거일)을 포함해 직접 표시했습니다.\n\n"
                  "스크립트: https://github.com/itmir913/teacher-utility-kit/blob/master/ml-playgrounds/datasets/tools/prepare_sources.py",
            "en": f"How much electricity all of Korea (mainland and Jeju) used each hour through {DEMAND_YEAR}, with the temperature at that hour. Predict demand from hour, weekday, holiday and temperature.\n"
                  "High in the day and low before dawn, lower at weekends and on holidays, higher in deep summer and deep winter for cooling and heating - guess first, then compare with the model.\n"
                  "Plot temperature against demand and a U shape appears, lowest in spring and autumn - a good reason to ask why one straight line does not fit.\n\n"
                  "Leave `date` out of training (it repeats across 24 rows a day) and use it to split train and test in time order.\n\n"
                  "## How this was built - two sources joined\n\n"
                  f"1. **Demand**: the {DEMAND_YEAR} file of Korea Power Exchange's \"Hourly national electricity demand\" on data.go.kr. The original is one row per day with hours 1-24 across 25 columns (365 rows); it was unfolded into 24 rows per day. Per the publisher, the figures are at the generation end and are provisional, for demand forecasting.\n"
                  "2. **Temperature**: the demand file has none. Hourly 2 m temperature for five cities - Seoul, Busan, Daegu, Gwangju and Daejeon - was taken from the Open-Meteo historical weather API and averaged hour by hour, so the figure is not just one city's weather.\n"
                  "3. **Aligning hours**: the original '1' is the energy used from 00:00 to 01:00; it gets the temperature at the end of that interval. Hour 24 gets the next day's 00:00 temperature.\n"
                  "4. **Calendar columns**: month and weekday come from the date. Holidays were marked by hand, including substitute holidays and the one-off holidays of 27 January and 3 June (presidential election).\n\n"
                  "Script: https://github.com/itmir913/teacher-utility-kit/blob/master/ml-playgrounds/datasets/tools/prepare_sources.py",
        },
        "notes": {
            "ko": ["원본 열 이름(날짜, 1시~24시)은 한국어였습니다. 세로로 펼치면서 영어 열 이름을 새로 붙였습니다.",
                   "원본은 365일이고 결측은 없었습니다. 파일 끝의 쉼표만 있는 줄은 뺐습니다."],
            "en": ["The original column names (date, hours 1-24) were in Korean; English names were given when the table was unfolded.",
                   "The original covers 365 days with no missing values; a trailing line of bare commas was dropped."],
        },
    })


# ── 신용카드 이상거래 (Sparkov 합성 자료) ──────────────────────────────────
FRAUD_RATIO = 4      # 사기 1건에 정상 몇 건을 붙일지 (사기가 모자라는 큰 파일은 예외). 원본의 사기 비율은 약 0.5%다.
FRAUD_SEED = 20261014
FRAUD_SIZES = [100_000, 48_000, 10_000]   # 교실 컴퓨터 사정에 맞춰 고르도록 세 크기를 함께 넣는다
FRAUD_ALL = 48_000   # 이 판은 행 수를 딱 맞추지 않고 사기를 전부 넣는다 (사기 × (1 + FRAUD_RATIO)행)


def fraud_file(n):
    return f"credit-card-fraud-{n // 1000}k.csv"

FRAUD_CATEGORIES = {
    "entertainment": "오락", "food_dining": "외식", "gas_transport": "주유·교통",
    "grocery_net": "식료품 (온라인)", "grocery_pos": "식료품 (매장)", "health_fitness": "건강·운동",
    "home": "생활용품", "kids_pets": "아동·반려동물", "misc_net": "기타 (온라인)", "misc_pos": "기타 (매장)",
    "personal_care": "미용·위생", "shopping_net": "쇼핑 (온라인)", "shopping_pos": "쇼핑 (매장)", "travel": "여행",
}


def credit_card_fraud():
    import datetime as dt, math, random
    frauds, normals, total = [], [], 0
    for name in ("fraudTrain.csv", "fraudTest.csv"):   # 둘은 기간만 다르다. 사기를 하나라도 더 남기려고 합친다.
        with open(os.path.join(RAW, "fraud", name), encoding="utf-8", newline="") as f:
            rd = csv.DictReader(f)
            for r in rd:
                total += 1
                (frauds if r["is_fraud"] == "1" else normals).append(r)
    # 크기별로 세 파일을 만든다. 작은 파일이 큰 파일의 부분집합이 되도록 한 번 섞은
    # 목록의 앞에서부터 자른다 - 크기만 다를 뿐 같은 자료다.
    rng = random.Random(FRAUD_SEED)
    rng.shuffle(frauds)
    rng.shuffle(normals)
    n_f = len(frauds)
    sizes = {}
    for rows_total in FRAUD_SIZES:
        if rows_total == FRAUD_ALL:
            sizes[rows_total] = (n_f, n_f * FRAUD_RATIO)
            continue
        nf = min(n_f, rows_total // (FRAUD_RATIO + 1))
        sizes[rows_total] = (nf, rows_total - nf)

    def km(a1, o1, a2, o2):
        a1, o1, a2, o2 = map(math.radians, (a1, o1, a2, o2))
        h = math.sin((a2 - a1) / 2) ** 2 + math.cos(a1) * math.cos(a2) * math.sin((o2 - o1) / 2) ** 2
        return 6371.0 * 2 * math.asin(math.sqrt(h))

    def row(r):
        t = dt.datetime.fromisoformat(r["trans_date_trans_time"])
        b = dt.date.fromisoformat(r["dob"])
        age = t.year - b.year - ((t.month, t.day) < (b.month, b.day))
        return [str(t.hour), WEEKDAYS[t.weekday()], r["category"], r["amt"], r["gender"], str(age),
                r["city_pop"],
                f"{km(float(r['lat']), float(r['long']), float(r['merch_lat']), float(r['merch_long'])):.1f}",
                r["is_fraud"]]

    header = ["hour", "weekday", "category", "amount (USD)", "gender", "age", "city population", "distance (km)", "is_fraud"]
    print(f"  credit-card-fraud: 원본 {total:,}행, 사기 {n_f:,}건 ({n_f / total * 100:.2f}%)")
    for rows_total, (nf, nn) in sizes.items():
        picked = frauds[:nf] + normals[:nn]
        random.Random(FRAUD_SEED + rows_total).shuffle(picked)
        write(fraud_file(rows_total), header, [row(r) for r in picked])
        print(f"    {fraud_file(rows_total)}: 사기 {nf:,} + 정상 {nn:,} = {nf + nn:,}행 ({nf / (nf + nn) * 100:.1f}%)")

    def size_table(lang):
        head = "| 파일 | 행 | 사기 | 정상 | 사기 비율 |" if lang == "ko" else "| File | Rows | Fraud | Normal | Fraud rate |"
        lines = [head, "|---|---|---|---|---|"]
        for rows_total, (nf, nn) in sizes.items():
            lines.append(f"| `{fraud_file(rows_total)}` | {nf + nn:,} | {nf:,} | {nn:,} | {nf / (nf + nn) * 100:.0f}% |")
        return "\n".join(lines)

    recipe({
        "id": "credit-card-fraud",
        "csv_name": [fraud_file(n) for n in FRAUD_SIZES],
        "source_url": "https://www.kaggle.com/datasets/kartik2112/fraud-detection",
        "license": "CC0: Public Domain (자료), MIT (생성 프로그램 https://github.com/namebrandon/Sparkov_Data_Generation)",
        "title": {"ko": "신용카드 이상거래", "en": "Credit Card Fraud"},
        "target": {"ko": "사기 여부 (is_fraud)", "en": "is_fraud"},
        "headers": [
            H("hour", "거래 시각 (hour)"),
            H("weekday", "요일 (weekday)"),
            H("category", "업종 (category)"),
            H("amount (USD)", "금액 (amount, 달러)"),
            H("gender", "카드 주인 성별 (gender)"),
            H("age", "카드 주인 나이 (age)"),
            H("city population", "사는 도시 인구 (city population)"),
            H("distance (km)", "집과 가게 사이 거리 (distance, km)"),
            H("is_fraud", "사기 여부 (is_fraud)"),
        ],
        "values": {
            "weekday": {"ko": {e: f"{k} ({e})" for e, k in zip(WEEKDAYS, WEEKDAYS_KO)}},
            "category": {"ko": {e: f"{k} ({e})" for e, k in FRAUD_CATEGORIES.items()}},
            "gender": {"ko": {"F": "여성 (F)", "M": "남성 (M)"}, "en": {"F": "Female (F)", "M": "Male (M)"}},
            "is_fraud": {"ko": {"0": "정상 (0)", "1": "사기 (1)"}, "en": {"0": "Normal (0)", "1": "Fraud (1)"}},
        },
        "body": {
            "ko": "미국의 카드 거래를 흉내 내어 만든 **합성 자료**입니다. 실제 사람의 거래가 아닙니다. 거래 시각·업종·금액·카드 주인의 나이 등으로 사기 거래를 가려냅니다.\n"
                  "늦은 밤에, 큰 금액이, 온라인 쇼핑에서 결제되면 수상하다는 짐작을 확인해 보세요. 로지스틱 회귀로 사기일 확률을 보거나, 정답 열을 빼고 군집으로 튀는 거래 무리가 생기는지 볼 수 있습니다.\n\n"
                  f"## 사기 비율을 일부러 높였습니다\n\n"
                  f"원본에서 사기는 {n_f / total * 100:.2f}%뿐입니다. 그대로 쓰면 \"전부 정상\"이라고만 답하는 모델이 정확도 99% 이상을 받고, 사기는 한 건도 잡지 못합니다. "
                  "그래서 사기 거래를 남기고 정상 거래를 무작위로 뽑아 섞었습니다. 교실 컴퓨터 사정에 맞춰 고를 수 있도록 크기가 다른 파일 세 개를 넣었습니다.\n\n"
                  + size_table("ko") + "\n\n"
                  f"사기 거래는 원본 전체를 통틀어 {n_f:,}건뿐이라, 10만 행 파일은 사기 비율이 다른 두 파일보다 낮습니다. "
                  "작은 파일은 큰 파일의 부분집합입니다 - 크기만 다를 뿐 같은 자료에서 잘라 낸 것입니다. 처음에는 가장 작은 파일로 시작하십시오.\n"
                  "실제 세상의 사기는 훨씬 드물다는 것을 기억하십시오. 이 비율의 차이 자체가 좋은 토론거리입니다.\n\n"
                  "## 이렇게 만들었습니다\n\n"
                  f"1. 원본의 학습용(fraudTrain.csv)과 테스트용(fraudTest.csv)을 합쳤습니다 (총 {total:,}건). 두 파일은 기간만 다릅니다.\n"
                  f"2. 사기와 정상 거래를 각각 고정된 난수 씨앗({FRAUD_SEED})으로 섞은 뒤, 위 표의 개수만큼 앞에서부터 가져와 다시 섞었습니다.\n"
                  "3. 거래 일시에서 시각(0~23)과 요일을 뽑았습니다.\n"
                  "4. 생년월일과 거래일로 거래 당시 나이를 계산했습니다.\n"
                  "5. 카드 주인 집과 가게의 위도·경도로 둘 사이 거리(km, 대원 거리)를 계산했습니다.\n"
                  "6. 이름·주소·카드번호·거래번호·직업 등 가짜 개인정보 열과 가게 이름, 원본 좌표는 뺐습니다. 지어낸 값이지만 실제 개인정보처럼 보이기 때문입니다.\n\n"
                  "거리 열은 원본 생성 방식상 가게 위치를 집 근처에 무작위로 흩어 놓았기 때문에 사기와 거의 관련이 없습니다. \"쓸모없는 특징\"을 찾아보는 연습으로 쓰십시오.\n\n"
                  "스크립트: https://github.com/itmir913/teacher-utility-kit/blob/master/ml-playgrounds/datasets/tools/prepare_sources.py",
            "en": "A **synthetic** set imitating US card transactions - no real person's purchases. Pick out fraud from the hour, merchant category, amount, the cardholder's age and so on.\n"
                  "Check the hunch that late-night, large, online purchases look suspicious. Use logistic regression to read off a fraud probability, or drop the label and cluster to see whether a group of odd transactions stands apart.\n\n"
                  "## The fraud rate was raised on purpose\n\n"
                  f"Only {n_f / total * 100:.2f}% of the original transactions are fraud. Used as is, a model that answers \"normal\" every time scores over 99% accuracy and catches nothing. "
                  "So frauds were kept and normal transactions drawn at random. Three files of different sizes are included, so you can pick one your classroom computers can handle.\n\n"
                  + size_table("en") + "\n\n"
                  f"There are only {n_f:,} frauds in the whole original, so the 100k file has a lower fraud rate than the other two. "
                  "Each smaller file is a subset of the larger one - the same data, cut shorter. Start with the smallest.\n"
                  "Remember that real-world fraud is far rarer - the gap itself is worth discussing.\n\n"
                  "## How this was built\n\n"
                  f"1. The original training file (fraudTrain.csv) and test file (fraudTest.csv) were combined ({total:,} transactions); they differ only in period.\n"
                  f"2. Frauds and normal transactions were each shuffled with a fixed random seed ({FRAUD_SEED}); each file takes the counts in the table above from the front and is shuffled again.\n"
                  "3. Hour (0-23) and weekday come from the transaction timestamp.\n"
                  "4. Age at the time of the transaction comes from the date of birth.\n"
                  "5. Distance (km, great-circle) is computed from the cardholder's and the merchant's latitude and longitude.\n"
                  "6. Fake personal columns - names, address, card number, transaction ID, job - plus merchant name and raw coordinates were dropped; invented, but they look like real personal data.\n\n"
                  "Because the generator scatters merchants randomly around the cardholder's home, distance has almost nothing to do with fraud. Use it to practise spotting a useless feature.\n\n"
                  "Script: https://github.com/itmir913/teacher-utility-kit/blob/master/ml-playgrounds/datasets/tools/prepare_sources.py",
        },
        "notes": {
            "ko": ["위 \"이렇게 만들었습니다\"에 적은 것이 전부입니다."],
            "en": ["Everything changed is listed under \"How this was built\" above."],
        },
    })


# ── 가위바위보 (이미지) ──────────────────────────────────────────────────
RPS_SIZE = 150       # 원본 300×300을 절반으로. 손 모양은 그대로 보이고 용량은 1/4 남짓이 된다.


def rock_paper_scissors():
    """이미지 자료는 build_dataset.py를 거치지 않는다. zip을 바로 만든다 (Pillow 필요)."""
    import io, zipfile
    from PIL import Image
    out_zip = os.path.join(CLEAN, "rock-paper-scissors.zip")
    counts = {}
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_STORED) as z:   # JPG는 이미 압축되어 있다
        for src, split in (("rps.zip", "training_set"), ("rps-test.zip", "test_set")):
            with zipfile.ZipFile(os.path.join(RAW, src)) as zin:
                for n in sorted(zin.namelist()):
                    parts = n.split("/")
                    if len(parts) != 3 or not parts[2].lower().endswith(".png"):
                        continue
                    label = parts[1]
                    assert label in ("rock", "paper", "scissors"), n
                    img = Image.open(io.BytesIO(zin.read(n))).convert("RGB").resize((RPS_SIZE, RPS_SIZE), Image.LANCZOS)
                    buf = io.BytesIO()
                    img.save(buf, "JPEG", quality=85)
                    z.writestr(f"{split}/{label}/{parts[2][:-4]}.jpg", buf.getvalue())
                    counts[(split, label)] = counts.get((split, label), 0) + 1
        z.writestr("README.md", f"""# Rock Paper Scissors

Source: https://laurencemoroney.com/datasets.html
License: CC BY 2.0 — Laurence Moroney (lmoroney@gmail.com / laurencemoroney.com)

Computer-generated (CGI) hands of different ages, skin tones and genders,
posed as rock, paper or scissors on a white background. Split into
training_set and test_set folders; the folder name is the label.
Choose clustering instead and those names are set aside.

## Changed from the original

- Resized from 300×300 PNG to {RPS_SIZE}×{RPS_SIZE} JPG (quality 85) to keep the download small.
- The training set (rps.zip) became training_set and the test set (rps-test-set.zip) became test_set.
""".encode("utf-8"))
    for k in sorted(counts):
        print(f"  {k[0]}/{k[1]}: {counts[k]}")
    print(f"  rock-paper-scissors.zip  {os.path.getsize(out_zip):,} bytes")


STEPS = {
    "galton-heights": galton_heights,
    "diabetes": diabetes,
    "titanic": titanic,
    "penguins": penguins,
    "zoo": zoo,
    "mushroom": mushroom,
    "power-plant": power_plant,
    "real-estate": real_estate,
    "beijing-pm25": beijing_pm25,
    "power-demand": power_demand,
    "credit-card-fraud": credit_card_fraud,
    "rock-paper-scissors": rock_paper_scissors,
}

if __name__ == "__main__":
    for step in sys.argv[4:] or STEPS:
        print(step)
        STEPS[step]()
